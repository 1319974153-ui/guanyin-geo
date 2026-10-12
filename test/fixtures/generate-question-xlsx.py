# Synthetic interoperability fixture; generated using openpyxl, not the importer.
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font
book = Workbook()
book.properties.creator = 'Guanyin test fixture'
book.properties.created = datetime(2026, 9, 29)
intro = book.active
intro.title = '说明'
intro['A1'] = '请选择客户问题工作表，第三行为表头。'
sheet = book.create_sheet('客户问题')
sheet['A1'] = '问题导入互操作测试'
sheet.append([])
for col, value in enumerate(['问题正文', '主题', '用户意图', '标签', '产品线', '备注'], 1):
    sheet.cell(3, col, value).font = Font(bold=True)
values = ['产品如何使用？', '上手', '使用', '教程,入门', '品牌通用', '普通文字']
for col, value in enumerate(values, 1):
    sheet.cell(4, col, value)
sheet['A6'] = '如何迁移\n已有资料？'
sheet['C6'] = '了解'
sheet['F6'] = '=HYPERLINK("https://example.invalid", "不执行")'
sheet.row_dimensions[6].hidden = True
sheet['A7'] = '=CONCAT("公式", "问题")'
sheet['C7'] = '了解'
sheet['A8'] = '#VALUE!'
sheet['A8'].data_type = 'e'
sheet['A10'] = '如何核对官方网站？'
sheet['F10'] = 'https://example.invalid/info'
sheet['F10'].hyperlink = 'https://example.invalid/info'
sheet.merge_cells('G20:H20')
private = book.create_sheet('隐藏问题')
private.sheet_state = 'hidden'
private.append(['问题正文'])
private.append(['隐藏表中的测试问题'])
book.save(Path(__file__).with_name('questions-openpyxl.xlsx'))
